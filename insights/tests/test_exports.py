"""BI exports (Phase 8, spec §28–29): anonymous, aggregated, honest about simulated data."""

from __future__ import annotations

import csv
import zipfile
from datetime import datetime, timedelta, timezone

import pytest

from polybot_insights.exports.build import ExportError, build_exports, learner_key
from polybot_insights.model.build import build_learner_model
from polybot_insights.personalization.engine import run_personalization
from polybot_insights.pipeline.run import run_pipeline
from polybot_insights.simulator import Simulator

SALT = "test-salt-0123456789abcdef"
SHEETS = ["Summary", "Skill Performance", "Content Performance", "Engagement", "Personalization Outcomes", "Raw Analytics Data"]


@pytest.fixture
def exported(pdb, tmp_path):
    sim = Simulator(pdb, seed=41, invalid_rate=0.0, end=datetime.now(timezone.utc) - timedelta(hours=1))
    sim.run(learners=5, days=21)
    run_pipeline(pdb, "test")
    build_learner_model(pdb, backfill_days=3, trigger_source="test")
    run_personalization(pdb, "test")
    result = build_exports(pdb, tmp_path, SALT, "test")
    return sim, result, tmp_path


def read_csv(path):
    with path.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def test_all_files_are_written_and_tracked(exported, pdb):
    _, result, out = exported
    assert result["status"] == "succeeded"
    assert sorted(result["files"]) == sorted(["polybot_analytics.csv", "engagement_weekly.csv", "content_performance.csv",
                                              "personalization_outcomes.csv", "polybot_analytics.xlsx"])
    for name in result["files"]:
        assert (out / name).stat().st_size > 0
    run = pdb.execute("SELECT status, rows_processed FROM ops.pipeline_runs WHERE run_id = %s", (result["run_id"],)).fetchone()
    assert run == ("succeeded", result["rows"])


def test_the_main_tableau_source_is_tidy_and_labelled(exported):
    sim, _, out = exported
    rows = read_csv(out / "polybot_analytics.csv")
    assert rows and {"learner_key", "is_simulated", "language", "date", "skill", "skill_group", "mastery_score",
                     "confidence_label"} <= rows[0].keys()
    assert {r["is_simulated"] for r in rows} == {"True"}
    assert len({r["date"] for r in rows}) == 4, "today plus 3 backfilled days"
    expected = {learner_key(SALT.encode(), l.learner_id) for l in sim.learners}
    assert {r["learner_key"] for r in rows} <= expected


def test_no_learner_id_or_vocabulary_term_leaves_the_database(exported, pdb):
    _, _, out = exported
    ids = [str(r[0]) for r in pdb.execute("SELECT learner_id FROM core.learners")]
    words = [r[0] for r in pdb.execute("SELECT DISTINCT vocab_item_id FROM core.learner_events WHERE vocab_item_id IS NOT NULL")]
    phrases = [r[0] for r in pdb.execute("SELECT lesson_id FROM core.lessons WHERE owner_learner_id IS NOT NULL")]
    texts = [p.read_text(encoding="utf-8") for p in out.glob("*.csv")]
    with zipfile.ZipFile(out / "polybot_analytics.xlsx") as z:  # check the workbook's own XML too
        texts += [z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml")]
    blob = "\n".join(texts)
    assert ids and not [i for i in ids if i in blob]
    assert words and not [w for w in words if w in blob]
    assert phrases and not [p for p in phrases if p in blob]


def test_keys_depend_on_the_secret_salt():
    learner = "6c0e2f6b-5b9e-4d8a-8f2a-0d4f7b1e9c33"
    a, b = learner_key(b"salt-one-xxxxxxxx", learner), learner_key(b"salt-two-xxxxxxxx", learner)
    assert a != b and len(a) == 16 and learner_key(b"salt-one-xxxxxxxx", learner) == a


def test_the_workbook_has_the_six_tabs_as_excel_tables(exported):
    from openpyxl import load_workbook
    _, _, out = exported
    wb = load_workbook(out / "polybot_analytics.xlsx")
    assert wb.sheetnames == SHEETS
    for name in SHEETS[1:]:
        assert len(wb[name].tables) == 1, f"{name} should be an Excel Table for PivotTables"
    summary = wb["Summary"]
    values = [c.value for row in summary.iter_rows() for c in row if c.value is not None]
    assert any(isinstance(v, str) and v.startswith("Contains 5 simulated learner") for v in values)
    assert "=SUM(PersonalizationOutcomes[accepted])" in values
    assert "=ROWS(RawAnalyticsData[learner_key])" in values


def test_content_performance_uses_catalog_content_only(exported):
    _, _, out = exported
    rows = read_csv(out / "content_performance.csv")
    kinds = {r["content_kind"] for r in rows}
    assert kinds == {"lesson", "exercise_type"}
    assert not [r for r in rows if r["content_id"].startswith(("phrase:", "media:es:sim"))]
    recall = [r for r in rows if r["content_type"] == "vocab_recall_choice"]
    assert recall and all(0 <= float(r["first_try_accuracy"]) <= 1 for r in recall)


def test_exports_refuse_to_run_without_a_real_salt(pdb, tmp_path):
    for bad in ("", "REPLACE_WITH_A_LONG_RANDOM_STRING", "short"):
        with pytest.raises(ExportError, match="POLYBOT_EXPORT_SALT"):
            build_exports(pdb, tmp_path, bad, "test")
    assert not list(tmp_path.iterdir())
